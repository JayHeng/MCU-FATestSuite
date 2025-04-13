#! /usr/bin/env python
# -*- coding: utf-8 -*-

# Copyright 2024 NXP
# All rights reserved.
# 
# SPDX-License-Identifier: BSD-3-Clause

import sys
import os
import time
import bincopy
from openpyxl import Workbook
from . import rundef
from . import debugger_utils
from . import debugger_pyocd
from . import debugger_pylink
import boot
sys.path.append(os.path.abspath(".."))
from ui import uicore
from ui import uidef
from ui import uilang
from boot import target

def execfile(filepath, globals=None, locals=None):
    if globals is None:
        globals = {}
    globals.update({
        "__file__": filepath,
        "__name__": "__main__",
    })
    with open(filepath, 'rb') as file:
        exec(compile(file.read(), filepath, 'exec'), globals, locals)


def createTarget(device, exeBinRoot):
    cpu = device
    targetBaseDir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'targets', cpu)

    # Check for existing target directory.
    if not os.path.isdir(targetBaseDir):
        targetBaseDir = os.path.join(os.path.dirname(exeBinRoot), 'src', 'targets', cpu)
        if not os.path.isdir(targetBaseDir):
            raise ValueError("Missing target directory at path %s" % targetBaseDir)

    targetConfigFile = os.path.join(targetBaseDir, 'fattargetconfig.py')

    # Check for config file existence.
    if not os.path.isfile(targetConfigFile):
        raise RuntimeError("Missing target config file at path %s" % targetConfigFile)

    # Build locals dict by copying our locals and adjusting file path and name.
    targetConfig = locals().copy()
    targetConfig['__file__'] = targetConfigFile
    targetConfig['__name__'] = 'fattargetconfig'

    # Execute the target config script.
    execfile(targetConfigFile, globals(), targetConfig)

    # Create the target object.
    tgt = target.Target(**targetConfig)

    return tgt, targetBaseDir

##
# @brief
class faTesterRun(uicore.faTesterUi):

    def __init__(self, parent):
        uicore.faTesterUi.__init__(self, parent)
        self.tgt = None
        self.cpuDir = None
        self.createMcuTarget()

    def createMcuTarget( self ):
        self.tgt, self.cpuDir = createTarget(self.mcuDevice, self.exeBinRoot)

    def task_receiveUartData( self ):
        while True:
            if self.enableUartToRecvData:
                if self.serialPort.isOpen():
                    num = self.serialPort.inWaiting()
                    if num != 0:
                        data = self.serialPort.read(num)
                        try:
                            string = ''
                            for i in range(len(data)):
                                string += chr(data[i])
                            #string = data.decode()
                            self.recvPrintBuf += string
                            self.appendContentOnMainPrintWin(string)
                        except:
                            pass
            time.sleep(self.tgt.uartRecvInterval)

    def refreshJlinkSN( self ):
        if self.testLoader == uidef.kTestLoader_Jlink:
            self._pyocd = debugger_pyocd.PyocdDebugger()
            jlinkUid = self._pyocd.getJlinkUid()
            self.adjustLoaderSN(jlinkUid)

    def findTestCases( self ):
        self.resetTestResult(False)
        caseTestResultMsg = ""
        fwAppFiles = []
        fwAppNames = []
        fwFolderPath = os.path.join(self.exeTopRoot, 'src', 'targets', self.tgt.cpu, self.mcuBoard)
        files = os.listdir(fwFolderPath)
        for file in files:
            filename, filetype = os.path.splitext(file)
            if filetype == '.srec' or filetype == '.s19':
                fwAppFiles.append(os.path.join(fwFolderPath, file))
                caseTestResultMsg += "( TBD ) -- " + filename + "\n"
                fwAppNames.append(filename)
        self.fwAppFiles = fwAppFiles[:]
        self.fwAppNames = fwAppNames[:]
        if len(fwAppFiles) == 0:
            self.showInfoMessage('App Error', 'Cannot find any test case files (.srec/.s19)')
        else:
            self.appendContentOnMainResWin(caseTestResultMsg)
            self.setButtonProperty("detectTestCases", uidef.kButtonColor_Green)

    def _getVal32FromByteArray( self, binarray, offset=0):
        val32Vaule = ((binarray[3+offset]<<24) + (binarray[2+offset]<<16) + (binarray[1+offset]<<8) + binarray[0+offset])
        return val32Vaule

    def _getAppDelayTime( self, loc ):
        delay = 0
        magicLen = len(self.tgt.fatLogDelay)
        while True:
            if len(self.recvPrintBuf) > loc + magicLen:
                res = self.recvPrintBuf[loc:loc+magicLen+1].find(self.tgt.fatLogDelay)
                if (res != -1):
                    loc = loc + res + len(self.tgt.fatLogDelay)
                    time = ''
                    while len(self.recvPrintBuf) > loc:
                        if self.recvPrintBuf[loc] == 's':
                            delay = int(time)
                            break
                        else:
                            time += self.recvPrintBuf[loc]
                            loc = loc + 1
                else:
                    break
        #print('delay time = ' + str(delay))
        return delay

    def _flushTestResultLog (self, log ):
        self.caseResultLog += log
        self.appendContentOnMainResWin(log)

    def _loadTestCases( self ):
        if not os.path.isfile(self.loaderExe):
            self.showInfoMessage('Loader Error', 'You need to set Loader EXE first.')
            return
        self.resetTestResult(False)
        appLen = len(self.fwAppFiles)
        if appLen == 0:
            self.showInfoMessage('Flow Error', 'You need to detect test cases first.')
            return 
        self.setButtonProperty("runTestCases", uidef.kButtonColor_Yellow)
        self.recvPrintBuf = ""
        self.caseResultLog = ""
        fwAppResults = []
        if self.serialPort.isOpen():
            self.serialPort.reset_input_buffer()
        else:
            self.showInfoMessage('Flow Error', 'Com Port is not opened.')
            return 
        jlinkcmdFolderPath = os.path.join(self.exeTopRoot, 'src', 'run', 'debuggers', 'jlink')
        #print('Creating JLink debugger object...')
        if self.testLoader == uidef.kTestLoader_Jlink:
            self._debugger = debugger_utils.createDebugger(debugger_utils.kDebuggerType_JLink, self.tgt.jlinkDevice, self.tgt.jlinkInterface, self.tgt.jlinkSpeedInkHz, self.loaderExe, jlinkcmdFolderPath)
        else:
            return
        self._debugger.open()
        #print('Created JLink debugger object\r\n')
        lastBeg = 0
        for appIdx in range(appLen):
            self.setButtonProperty("runTestCases", None, 'Running Test Case ' + str(appIdx+1) + '/' + str(appLen))
            self.appendContentOnMainPrintWin('---------Case ' + str(appIdx+1) + '/' + str(appLen) + '----------\n')
            srecObj = bincopy.BinFile(str(self.fwAppFiles[appIdx]))
            filepath, file = os.path.split(self.fwAppFiles[appIdx])
            filename, filetype = os.path.splitext(file)
            startAddress = srecObj.minimum_address
            initialAppBytes = srecObj.as_binary(startAddress, startAddress + 8)
            sp = self._getVal32FromByteArray(initialAppBytes[0:4])
            pc = self._getVal32FromByteArray(initialAppBytes[4:8])
            appIsLoaded = False
            loadAppRetryCount = 0
            while (not appIsLoaded):
                #print('Loading app ' + self.fwAppFiles[appIdx] + ' via JLink debugger...')
                self._debugger.JumpToApp(self.fwAppFiles[appIdx], sp, pc, None)
                #print('Loaded app via JLink debugger\r\n')
                deltaTimeStart_load = time.perf_counter()
                while True:
                    res0 = self.recvPrintBuf.find(self.tgt.fatLogStart, lastBeg)
                    ##############################################################
                    if (res0 != -1):
                        deltaTimeStart_check = time.perf_counter()
                        appIsLoaded = True
                        delayTimeApp = self._getAppDelayTime(res0 + len(self.tgt.fatLogStart))
                        lastBeg = res0
                        while True:
                            res1 = self.recvPrintBuf.find(self.tgt.fatLogPass, lastBeg)
                            res2 = self.recvPrintBuf.find(self.tgt.fatLogFail, lastBeg)
                            if (res1 != -1):
                                lastBeg = res1
                                self._flushTestResultLog('( ' + rundef.kTestResult_RunPass + ' ) ' + filename)
                                fwAppResults.append(rundef.kTestResult_RunPass)
                                if delayTimeApp != 0:
                                    self._flushTestResultLog(', <case delay ' + str(delayTimeApp) + 's>\n')
                                    deltaTimeAppStart = time.perf_counter()
                                    deltaTime_app = time.perf_counter() - deltaTimeAppStart
                                    while (deltaTime_app < delayTimeApp):
                                        deltaTime_app = time.perf_counter() - deltaTimeAppStart
                                        time.sleep(1)
                                else:
                                    self._flushTestResultLog('\n')
                                break
                            if (res2 != -1):
                                lastBeg = res2
                                self._flushTestResultLog('( ' + rundef.kTestResult_RunFail + ' ) ' + filename + '\n')
                                fwAppResults.append(rundef.kTestResult_RunFail)
                                break
                            deltaTime_check = time.perf_counter() - deltaTimeStart_check
                            if (deltaTime_check > self.tgt.waitAppTimeout):
                                self._flushTestResultLog('( ' + rundef.kTestResult_RunTimeout + ' ) ' + filename + '\n')
                                fwAppResults.append(rundef.kTestResult_RunTimeout)
                                time.sleep(1)
                                break
                            time.sleep(0.5)
                        break
                    ##############################################################
                    #status, res0 = self._debugger.readMem32(self.tgt.fatRegAddr)
                    if False: #status and ((res0 & 0xFF) == self.tgt.fatRegStart):
                        appIsLoaded = True
                        while True:
                            status, resx = self._debugger.readMem32(self.tgt.fatRegAddr)
                            if status:
                                resx = resx >> 24
                                if resx == self.tgt.fatRegPass:
                                    self._flushTestResultLog('( ' + rundef.kTestResult_RunPass + ' ) ' + filename + '\n')
                                    fwAppResults.append(rundef.kTestResult_RunPass)
                                    break
                                elif resx == self.tgt.fatRegFail:
                                    self._flushTestResultLog('( ' + rundef.kTestResult_RunFail + ' ) ' + filename + '\n')
                                    fwAppResults.append(rundef.kTestResult_RunFail)
                                    break
                            time.sleep(0.5)
                        break
                    ##############################################################
                    deltaTime_load = time.perf_counter() - deltaTimeStart_load
                    if (deltaTime_load > self.tgt.loadAppTimeout):
                        time.sleep(1)
                        loadAppRetryCount += 1
                        if loadAppRetryCount > self.tgt.loadAppRetryCount:
                            appIsLoaded = True
                            self._flushTestResultLog('( ' + rundef.kTestResult_LoadFail + ' ) ' + filename + '\n')
                            fwAppResults.append(rundef.kTestResult_LoadFail)
                        break
                    time.sleep(0.5)
                    ##############################################################
        self.setButtonProperty("runTestCases", uidef.kButtonColor_White, 'Run Test Cases')
        self.fwAppResults = fwAppResults[:]
        #self.flushContentOnMainPrintWin()

    def _areAllTestCasesPassed(self):
        for i in range(len(self.fwAppResults)):
            if self.fwAppResults[i] != rundef.kTestResult_RunPass:
                return False
        return True

    def _saveTestResultToText( self, reportName ):
        self.updateBoardSN()
        resFilename = reportName + '.txt'
        with open(resFilename, 'w+') as fileObj:
            fileObj.write("\r\n-----------case result log--------------\r\n")
            fileObj.write(self.caseResultLog)
            fileObj.write("\r\n-----------case print log---------------\r\n")
            fileObj.write(self.recvPrintBuf)
            fileObj.close()

    def _saveTestResultToExcel( self, reportName ):
        self.updateBoardSN()
        resFilename = reportName + '.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Results"
        ws['A1'] = 'Name'
        ws['B1'] = 'Result'
        for i in range(len(self.fwAppNames)):
            ws.cell(row=2 + i, column=1).value = self.fwAppNames[i]
            ws.cell(row=2 + i, column=2).value = self.fwAppResults[i]
        wb.save(resFilename)
        wb.close()

    def task_loadTestCases( self ):
        while True:
            if self.isLoadTestCasesTaskPending:
                self._loadTestCases()
                finalResult = 'FAIL'
                if self._areAllTestCasesPassed():
                    finalResult = 'PASS'
                reportName = os.path.join(self.exeTopRoot, 'report', finalResult + "_test_report_" + self.mcuDevice + "_" + self.mcuBoard + "_" + self.boardSN + "_" + time.strftime('%Y-%m-%d_%H.%M.%S',time.localtime(time.time())))
                self._saveTestResultToText(reportName)
                self._saveTestResultToExcel(reportName)
                self._flushTestResultLog('\r\nDONE')
                self.isLoadTestCasesTaskPending = False
            time.sleep(1)
